#!/usr/bin/env python3
import configparser
import json
import os
import sys
import time
from datetime import datetime, timezone
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

config = configparser.ConfigParser()
config.read(os.path.join(os.path.dirname(os.path.abspath(__file__)), "mist_org_audit_log.ini"))

API_BASE = config.get("mist", "api_base")
ORG_ID = config.get("mist", "org_id")
API_TOKEN = config.get("mist", "api_token")
# Date range for the audit trail pull. Format: dd/mm/yyyy
START_DATE = config.get("dates", "start_date")
END_DATE = config.get("dates", "end_date")
OUTPUT_DIR = os.path.expanduser(config.get("output", "directory"))

parsed = urlparse(API_BASE)
API_HOST = f"{parsed.scheme}://{parsed.netloc}"


def parse_ddmmyyyy(date_str, end_of_day=False):
    dt = datetime.strptime(date_str, "%d/%m/%Y")
    if end_of_day:
        dt = dt.replace(hour=23, minute=59, second=59)
    return int(dt.timestamp())


def format_eta(seconds):
    if seconds < 60:
        return f"{int(seconds)}s"
    m, s = divmod(int(seconds), 60)
    return f"{m}m {s}s"


def request_with_retry(url, headers, max_retries=5):
    for attempt in range(1, max_retries + 1):
        resp = requests.get(url, headers=headers)
        if resp.status_code == 429:
            retry_after = int(resp.headers.get("Retry-After", 30))
            print(f"\n  Rate limited - waiting {retry_after}s (attempt {attempt}/{max_retries})...")
            time.sleep(retry_after)
            continue
        resp.raise_for_status()
        return resp
    resp.raise_for_status()
    return resp


def fetch_org_info(token):
    headers = {"Authorization": f"Token {token}"}
    resp = request_with_retry(f"{API_BASE}/orgs/{ORG_ID}", headers)
    return resp.json()


def fetch_site_names(token):
    headers = {"Authorization": f"Token {token}"}
    site_map = {}
    page = 1
    api_calls = 0
    while True:
        url = f"{API_BASE}/orgs/{ORG_ID}/sites?limit=1000&page={page}"
        resp = request_with_retry(url, headers)
        api_calls += 1
        data = resp.json()
        if not data:
            break
        for site in data:
            site_map[site["id"]] = site.get("name", site["id"])
        if len(data) < 1000:
            break
        page += 1
    return site_map, api_calls


def fetch_audit_logs(token, start, end):
    headers = {"Authorization": f"Token {token}"}
    all_logs = []
    api_calls = 0
    page_times = []
    total = None
    url = f"{API_BASE}/orgs/{ORG_ID}/logs?limit=1000&start={start}&end={end}&sort=-timestamp"
    page = 0
    empty_first_page_retried = False
    while url:
        page += 1
        t0 = time.time()
        resp = request_with_retry(url, headers)
        page_time = time.time() - t0
        page_times.append(page_time)
        api_calls += 1
        data = resp.json()
        if total is None:
            total = data.get("total", 0)
        results = data.get("results", [])
        if not results:
            # Mist's audit-log search occasionally returns an empty first
            # page for a large historical window even when data exists -
            # confirmed by re-running the identical request moments later
            # and getting real results. Give it one retry before giving up.
            if page == 1 and not empty_first_page_retried:
                empty_first_page_retried = True
                page -= 1
                page_times.pop()
                total = None
                print("\r  First page came back empty - retrying once in case it was a transient hiccup...")
                time.sleep(5)
                continue
            break
        all_logs.extend(results)
        avg_time = sum(page_times) / len(page_times)
        pct = min(99, len(all_logs) / max(total, 1) * 100) if total else 0
        remaining = max(0, total - len(all_logs)) / 1000 * avg_time
        sys.stdout.write(f"\r  Page {page} | {len(all_logs)}/{total} entries ({pct:.0f}%) | {avg_time:.1f}s/page | ETA: {format_eta(remaining)}   ")
        sys.stdout.flush()
        next_path = data.get("next")
        if next_path:
            url = f"{API_HOST}{next_path}"
        else:
            break
    print(f"\r  Audit log fetch complete - {len(all_logs)} entries in {format_eta(sum(page_times))}                    ")
    return all_logs, api_calls


def style_header(ws, headers, fill_color):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)


def main():
    if API_TOKEN in ("", "YOUR_API_TOKEN_HERE"):
        print("Set api_token in mist_org_audit_log.ini before running.")
        return

    start = parse_ddmmyyyy(START_DATE)
    end = parse_ddmmyyyy(END_DATE, end_of_day=True)
    if start >= end:
        print("START_DATE must be before END_DATE.")
        return

    print("Fetching organisation info...")
    org_info = fetch_org_info(API_TOKEN)
    org_name = org_info.get("name", "Unknown")

    print(f"\n{'='*55}")
    print(f"  Mist Org Audit Log")
    print(f"{'='*55}")
    print(f"  Organisation:  {org_name}")
    print(f"  Org ID:        {ORG_ID}")
    print(f"  Date range:    {START_DATE} - {END_DATE}")
    print(f"{'='*55}")

    confirm = input("\n  Proceed? (y/n): ").strip().lower()
    if confirm != "y":
        print("  Aborted.")
        return

    start_time = time.time()

    print("\nFetching audit log entries from Mist API...")
    logs, log_api_calls = fetch_audit_logs(API_TOKEN, start, end)
    print(f"Total entries fetched: {len(logs)}")

    site_ids = set(l.get("site_id") for l in logs if l.get("site_id"))
    site_map = {}
    site_api_calls = 0
    if site_ids:
        print(f"Fetching names for {len(site_ids)} sites...")
        site_map, site_api_calls = fetch_site_names(API_TOKEN)

    print("Building Excel spreadsheet...")

    headers = [
        "Timestamp (UTC)", "Timestamp (Epoch)", "Admin Name", "Admin ID",
        "Site Name", "Site ID", "For Site", "Message", "Source IP",
        "Before", "After", "Log ID",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Log"
    style_header(ws, headers, "37474F")

    for r, log in enumerate(logs, 2):
        ts = log.get("timestamp")
        ts_str = datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC") if ts else ""
        site_id = log.get("site_id", "") or ""
        before = log.get("before")
        after = log.get("after")
        row = [
            ts_str,
            ts,
            log.get("admin_name", ""),
            log.get("admin_id", ""),
            site_map.get(site_id, ""),
            site_id,
            log.get("for_site", ""),
            log.get("message", ""),
            log.get("src_ip", ""),
            json.dumps(before) if before else "",
            json.dumps(after) if after else "",
            log.get("id", ""),
        ]
        for c, value in enumerate(row, 1):
            ws.cell(row=r, column=c, value=value)

    auto_width(ws)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    safe_org_name = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in org_name).strip().replace(" ", "_")
    safe_start = START_DATE.replace("/", "-")
    safe_end = END_DATE.replace("/", "-")
    filename = f"Mist_Org_Audit_Log_{safe_org_name}_{safe_start}_to_{safe_end}_{timestamp}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    wb.save(filepath)

    total_api = log_api_calls + site_api_calls
    elapsed = time.time() - start_time
    print(f"\n{'='*55}")
    print(f"  Mist Org Audit Log Summary - {org_name}")
    print(f"{'='*55}")
    print(f"  Entries fetched:      {len(logs)}")
    print(f"  Unique sites seen:    {len(site_ids)}")
    print(f"  API calls - logs:     {log_api_calls}")
    print(f"  API calls - sites:    {site_api_calls}")
    print(f"  Total API calls:      {total_api}")
    print(f"  Total elapsed time:   {format_eta(elapsed)}")
    print(f"{'='*55}")
    print(f"  Report saved: {filepath}")


if __name__ == "__main__":
    main()
